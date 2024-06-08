#!/usr/bin/env python
# coding: utf-8

# In[ ]:


import streamlit as st
from translate_codebase import (load_dictionaries, translate_with_google, translate_word, 
                                translate_excel, translate_pptx, translate_pdf, save_translated_file)
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from PyPDF2 import PdfFileReader
import io
import base64

def get_download_link(file_path, file_name, file_label):
    with open(file_path, "rb") as f:
        bytes_data = f.read()
    b64 = base64.b64encode(bytes_data).decode()
    return f'<a href="data:file/txt;base64,{b64}" download="{file_name}">{file_label}</a>'

translation_dict = {}

st.title("AI Translation Tool")

uploaded_csv_files = st.file_uploader("Upload CSV files", type="csv", accept_multiple_files=True)
if uploaded_csv_files:
    translation_dict = load_dictionaries(uploaded_csv_files)
    st.write("Dictionaries loaded successfully.")

text_to_translate = st.text_area("Enter text to translate")
target_language = st.selectbox("Select target language", ["en", "id"])

if st.button("Translate Text"):
    if text_to_translate:
        translated_text = translate_with_google(text_to_translate, target_language, translation_dict)
        st.write("Translated Text:")
        st.write(translated_text)
    else:
        st.error("Please enter text to translate.")

uploaded_file = st.file_uploader("Upload a document", type=["docx", "xlsx", "pptx", "pdf"])
if uploaded_file:
    file_type = uploaded_file.name.split('.')[-1]
    file_name = f"translated_document.{file_type}"
    if file_type == "docx":
        doc = Document(uploaded_file)
        translated_doc = translate_word(doc, target_language, translation_dict)
        translated_doc.save(file_name)
        st.success(f"Translated Word document saved as '{file_name}'")
    elif file_type == "xlsx":
        wb = load_workbook(uploaded_file)
        translated_wb = translate_excel(wb, target_language, translation_dict)
        translated_wb.save(file_name)
        st.success(f"Translated Excel spreadsheet saved as '{file_name}'")
    elif file_type == "pptx":
        prs = Presentation(uploaded_file)
        translated_prs = translate_pptx(prs, target_language, translation_dict)
        translated_prs.save(file_name)
        st.success(f"Translated PowerPoint presentation saved as '{file_name}'")
    elif file_type == "pdf":
        pdf_reader = PdfFileReader(uploaded_file)
        translated_pdf_writer = translate_pdf(pdf_reader, target_language, translation_dict)
        with open(file_name, "wb") as f:
            translated_pdf_writer.write(f)
        st.success(f"Translated PDF saved as '{file_name}'")

    download_link = get_download_link(file_name, file_name, 'Download Translated Document')
    st.markdown(download_link, unsafe_allow_html=True)
